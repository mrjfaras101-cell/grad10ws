import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_rtl(table):
    tblPr = table._tbl.tblPr
    bidi = OxmlElement('w:bidiVisual')
    tblPr.append(bidi)


def set_table_borders(table, color="CBD5E1", sz="6", val="single"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        borders.append(border)
    tblPr.append(borders)


def add_rtl_paragraph(doc, text="", bold=False, size=11, color=None, space_after=6, space_before=0, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = doc.add_paragraph()
    p.alignment = align
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.font.name = 'Arial'
        run.font.size = Pt(size)
        rPr = run._r.get_or_add_rPr()
        rtl = OxmlElement('w:rtl')
        rPr.append(rtl)
        if color:
            run.font.color.rgb = RGBColor(*color)
    return p


def add_rtl_run(p, text, bold=False, size=11, color=None, font_name='Arial', is_code=False):
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Consolas' if is_code else font_name
    run.font.size = Pt(size)
    if not is_code:
        rPr = run._r.get_or_add_rPr()
        rtl = OxmlElement('w:rtl')
        rPr.append(rtl)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def add_header_banner(doc, ws_num_title, project_title, concepts_text, theme_color_hex="1E3A8A", light_bg_hex="EFF6FF"):
    # Official School Header Table
    header_table = doc.add_table(rows=3, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(header_table)
    set_table_borders(header_table, color=theme_color_hex, sz="8")

    # Top row: School & Teacher info
    row0 = header_table.rows[0]
    c0_0 = row0.cells[0]
    c0_1 = row0.cells[1]
    set_cell_shading(c0_0, theme_color_hex)
    set_cell_shading(c0_1, theme_color_hex)
    set_cell_margins(c0_0, 120, 120, 180, 180)
    set_cell_margins(c0_1, 120, 120, 180, 180)

    p0 = c0_0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p0._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
    add_rtl_run(p0, "🏫 مدرسة صويلح الثانوية للبنين\n", bold=True, size=13, color=(255, 255, 255))
    add_rtl_run(p0, "مبحث الحاسوب والبرمجة - لغة بايثون (Python)", bold=False, size=10, color=(226, 232, 240))

    p1 = c0_1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p1._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
    add_rtl_run(p1, "👨‍🏫 إعداد المعلم: باسم النعيمي\n", bold=True, size=12, color=(255, 255, 255))
    add_rtl_run(p1, "الصف: الأول الثانوي الأكاديمي", bold=True, size=11, color=(253, 224, 71))

    # Row 1: Worksheet Title merged
    row1 = header_table.rows[1]
    cell_title = row1.cells[0].merge(row1.cells[1])
    set_cell_shading(cell_title, light_bg_hex)
    set_cell_margins(cell_title, 140, 140, 180, 180)
    pt = cell_title.paragraphs[0]
    pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pt._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
    add_rtl_run(pt, f"{ws_num_title}\n", bold=True, size=11, color=(30, 58, 138))
    add_rtl_run(pt, project_title, bold=True, size=15, color=(15, 23, 42))

    # Row 2: Student Info
    row2 = header_table.rows[2]
    c2_0, c2_1 = row2.cells[0], row2.cells[1]
    set_cell_shading(c2_0, "F8FAFC")
    set_cell_shading(c2_1, "F8FAFC")
    set_cell_margins(c2_0, 100, 100, 150, 150)
    set_cell_margins(c2_1, 100, 100, 150, 150)

    ps0 = c2_0.paragraphs[0]
    ps0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ps0._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
    add_rtl_run(ps0, "اسم الطالب: ....................................................", bold=True, size=10, color=(51, 65, 85))

    ps1 = c2_1.paragraphs[0]
    ps1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ps1._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
    add_rtl_run(ps1, "الشعبة: ( ...... )   |   التاريخ: .... / .... / 2026م", bold=True, size=10, color=(51, 65, 85))

    # Concepts bar
    p_concepts = add_rtl_paragraph(doc, space_before=8, space_after=8)
    add_rtl_run(p_concepts, "🎯 المفاهيم البرمجية المستهدفة: ", bold=True, size=10.5, color=(15, 118, 110))
    add_rtl_run(p_concepts, concepts_text, bold=False, size=10.5, color=(51, 65, 85))


def add_box_section(doc, title, body_lines, bg_hex="F0FDF4", border_hex="16A34A", title_color=(21, 128, 61)):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(tbl)
    set_table_borders(tbl, color=border_hex, sz="8")
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, bg_hex)
    set_cell_margins(cell, 120, 120, 180, 180)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
    p.paragraph_format.space_after = Pt(4)
    add_rtl_run(p, title, bold=True, size=11.5, color=title_color)

    for line in body_lines:
        pl = cell.add_paragraph()
        pl.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pl._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
        pl.paragraph_format.space_after = Pt(3)
        add_rtl_run(pl, line, bold=False, size=10.5, color=(30, 41, 59))


def add_code_box(doc, title, code_lines, is_solution=False):
    add_rtl_paragraph(doc, title, bold=True, size=11.5, color=(30, 58, 138), space_before=8, space_after=4)
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    border_col = "16A34A" if is_solution else "64748B"
    bg_col = "F0FDF4" if is_solution else "F8FAFC"
    set_table_borders(tbl, color=border_col, sz="8")
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, bg_col)
    set_cell_margins(cell, 120, 120, 180, 180)

    for idx, line in enumerate(code_lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        is_comment = line.strip().startswith("#")
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
        if is_comment:
            run.bold = True
            run.font.color.rgb = RGBColor(5, 150, 105)
        else:
            run.font.color.rgb = RGBColor(15, 23, 42)


def build_ws1(doc, include_solution=False):
    add_header_banner(
        doc,
        ws_num_title="ورقة عمل (1) - مشروع مصغر تطبيقي",
        project_title="🛒 مشروع: برنامج الكاشير الذكي في مقصف المدرسة",
        concepts_text="جملة الطباعة print() | جملة الإدخال input() | المعاملات الحسابية (* , -) | المقارنة (>=) | الجملة الشرطية if / else",
        theme_color_hex="1E3A8A",
        light_bg_hex="EFF6FF"
    )

    add_box_section(
        doc,
        title="💡 فكرة المشروع (تعلّم ذاتي مبسط):",
        body_lines=[
            "يريد مقصف مدرسة صويلح الثانوية للبنين تصميم برنامج بسيط يحسب فاتورة الطالب تلقائياً بالدينار الأردني.",
            "يطلب البرنامج من الطالب إدخال (اسمه)، و(سعر الوجبة الواحدة بالدينار)، و(عدد الوجبات) التي اشتراها، ثم يحسب المجموع الكلي.",
            "ولتشجيع الطلبة: إذا كانت قيمة الفاتورة (10 دنانير أردنية أو أكثر) يحصل الطالب على خصم فوري بقيمة (2 دينار أردني)!"
        ],
        bg_hex="EFF6FF",
        border_hex="2563EB",
        title_color=(30, 58, 138)
    )

    add_rtl_paragraph(doc, "🛠️ خطوات بناء المشروع (نفّذها بالترتيب):", bold=True, size=11.5, color=(15, 23, 42), space_before=8, space_after=4)
    steps = [
        "1) اطبع رسالة ترحيبية على الشاشة في بداية البرنامج باستخدام دالة print().",
        "2) اطلب من الطالب إدخال (اسمه) باستخدام دالة input() وخزّنه في متغير باسم name.",
        "3) اطلب إدخال (سعر الوجبة بالدينار) في متغير price باستخدام float(input(...))، و(عدد الوجبات) في متغير quantity باستخدام int(input(...)).",
        "4) احسب (المبلغ الإجمالي) في متغير total عن طريق ضرب السعر في عدد الوجبات (price * quantity) ثم اطبعه.",
        "5) استخدم الجملة الشرطية if للتحقق: إذا كان (total >= 10)، اطرح 2 من المجموع واطبع رسالة تهنئة بالخصم، وإلا (else) اطبع أنه لا يوجد خصم، ثم اطبع المبلغ النهائي بالدينار الأردني."
    ]
    for s in steps:
        add_rtl_paragraph(doc, s, size=10.5, color=(30, 41, 59), space_after=3)

    add_box_section(
        doc,
        title="🔑 تلميحات ذكية لمساعدتك (Hints):",
        body_lines=[
            "• تلميح (1): دالة input() تقرأ المدخلات كنصوص دائماً؛ لذلك استخدم float() مع السعر و int() مع العدد.",
            "• تلميح (2): لدمج نص مع متغير داخل جملة الطباعة افصل بينهما بفاصلة مثل: print('مرحباً يا', name).",
            "• تلميح (3): لا تنسَ وضع النقطتين الرأسيتين (:) في نهاية سطر if وسطر else، وترك مسافة بادئة (Tab) للأسطر التي بداخلها.",
            "• مثال للمخرجات المتوقعة: عند إدخال الاسم (أحمد)، سعر الوجبة (4) دنانير، والعدد (3) -> المجموع قبل الخصم: 12.0 دينار -> بعد خصم 2 دينار يصبح المبلغ النهائي: 10.0 دينار أردني."
        ],
        bg_hex="FFFBEB",
        border_hex="D97706",
        title_color=(180, 83, 9)
    )

    skeleton_1 = [
        '# 1. طباعة رسالة الترحيب وقراءة بيانات الطالب',
        'print(..................................................)',
        'name = ................("أدخل اسمك: ")',
        'price = float(................("أدخل سعر الوجبة بالدينار الأردني: "))',
        'quantity = int(input("أدخل عدد الوجبات: "))',
        '',
        '# 2. العملية الحسابية لحساب المجموع بالدينار الأردني',
        'total = price ........ quantity',
        'print("المجموع قبل الخصم:", ..............., "دينار أردني")',
        '',
        '# 3. الجملة الشرطية للتحقق من استحقاق الخصم (10 دنانير فأكثر)',
        'if total ........ 10:',
        '    total = total - ........',
        '    print("مبروك يا", name, "حصلت على خصم بقيمة 2 دينار أردني!")',
        '............:',
        '    print("شكراً لك يا", name, "لا يوجد خصم لهذه الفاتورة.")',
        '',
        'print("المبلغ النهائي المطلوب دفعه:", total, "دينار أردني")'
    ]
    add_code_box(doc, "💻 مساحة كتابة الكود (أكمل الفراغات البرمجية التالية):", skeleton_1, is_solution=False)

    add_rtl_paragraph(
        doc,
        "🚀 تحدي المبدعين (إضافي): عدّل البرنامج ليطلب من الطالب إدخال (المبلغ النقدي المدفوع بالدينار)، ثم يحسب الباقي للطالب باستخدام عملية الطرح (-) ويطبعه!",
        bold=True, size=10, color=(109, 40, 217), space_before=6, space_after=4
    )

    if include_solution:
        doc.add_page_break()
        add_header_banner(
            doc,
            ws_num_title="نموذج الإجابة المعتمد - ورقة عمل (1)",
            project_title="✅ الحل النموذجي: برنامج الكاشير الذكي في مقصف المدرسة",
            concepts_text="إشراف المعلم: أ. باسم النعيمي | مدرسة صويلح الثانوية للبنين | الصف الأول الثانوي الأكاديمي",
            theme_color_hex="15803D",
            light_bg_hex="F0FDF4"
        )
        sol_1 = [
            '# 1. طباعة رسالة الترحيب وقراءة بيانات الطالب',
            'print("*** أهلاً بك في كاشير مقصف مدرسة صويلح الثانوية للبنين ***")',
            'name = input("أدخل اسمك: ")',
            'price = float(input("أدخل سعر الوجبة بالدينار الأردني: "))',
            'quantity = int(input("أدخل عدد الوجبات: "))',
            '',
            '# 2. العملية الحسابية لحساب المجموع بالدينار الأردني',
            'total = price * quantity',
            'print("المجموع قبل الخصم:", total, "دينار أردني")',
            '',
            '# 3. الجملة الشرطية للتحقق من استحقاق الخصم (10 دنانير فأكثر)',
            'if total >= 10:',
            '    total = total - 2',
            '    print("مبروك يا", name, "حصلت على خصم بقيمة 2 دينار أردني!")',
            'else:',
            '    print("شكراً لك يا", name, "لا يوجد خصم لهذه الفاتورة.")',
            '',
            'print("المبلغ النهائي المطلوب دفعه:", total, "دينار أردني")',
            '',
            '# --- حل تحدي المبدعين الإضافي (حساب الباقي) ---',
            'paid = float(input("أدخل المبلغ النقدي المدفوع بالدينار: "))',
            'change = paid - total',
            'print("الباقي لك:", change, "دينار أردني")'
        ]
        add_code_box(doc, "🗝️ الكود البرمجي الكامل مع حل التحدي الإضافي:", sol_1, is_solution=True)


def build_ws2(doc, include_solution=False):
    add_header_banner(
        doc,
        ws_num_title="ورقة عمل (2) - مشروع مصغر تطبيقي",
        project_title="🎓 مشروع: المرشد الأكاديمي وحاسبة المعدل الذكية",
        concepts_text="الأولويات الحسابية ( + ) / | الشروط المتعددة if / elif / else | معاملات المقارنة (>= , ==) | المعامل المنطقي (and)",
        theme_color_hex="0F766E",
        light_bg_hex="F0FDFA"
    )

    add_box_section(
        doc,
        title="💡 فكرة المشروع (تعلّم ذاتي مبسط):",
        body_lines=[
            "نريد بناء برنامج يساعد طالب الصف الأول الثانوي الأكاديمي في مدرسة صويلح الثانوية للبنين على معرفة نتيجته الأكاديمية.",
            "يطلب البرنامج من الطالب إدخال علاماته في (3 اختبارات)، وعدد (أيام غيابه)، ثم يحسب المعدل (المتوسط الحسابي).",
            "بعد ذلك يحدد التقدير اللفظي للطالب، ويقرر هل يستحق (شهادة التفوق والانضباط ومكافأة 25 ديناراً أردنياً) والتي تشترط أن يكون المعدل 90 فأكثر (و) عدد أيام الغياب يساوي 0."
        ],
        bg_hex="F0FDFA",
        border_hex="0D9488",
        title_color=(15, 118, 110)
    )

    add_rtl_paragraph(doc, "🛠️ خطوات بناء المشروع (نفّذها بالترتيب):", bold=True, size=11.5, color=(15, 23, 42), space_before=8, space_after=4)
    steps = [
        "1) اقرأ من الطالب علامات الاختبارات الثلاثة (g1, g2, g3) وعدد أيام الغياب (absence) كأعداد صحيحة باستخدام int(input(...)).",
        "2) احسب المعدل في متغير avg بجمع العلامات الثلاث بين قوسين وقسمتها على 3: (g1 + g2 + g3) / 3 ثم اطبعه.",
        "3) استخدم الجمل الشرطية المتعددة if / elif / else لطباعة التقدير: (>= 90 ممتاز) | (>= 75 جيد جداً) | (>= 50 ناجح) | (غير ذلك: يحتاج إلى تحسين).",
        "4) استخدم جملة if مع المعامل المنطقي and للتحقق مما إذا كان (avg >= 90) و (absence == 0) لطباعة رسالة منحه شهادة التفوق ومكافأة 25 ديناراً أردنياً!"
    ]
    for s in steps:
        add_rtl_paragraph(doc, s, size=10.5, color=(30, 41, 59), space_after=3)

    add_box_section(
        doc,
        title="🔑 تلميحات ذكية لمساعدتك (Hints):",
        body_lines=[
            "• تلميح (1) الأقواس مهمة!: يجب وضع عملية الجمع بين قوسين (g1 + g2 + g3) / 3 لأن القسمة لها أولوية أعلى من الجمع في البايثون.",
            "• تلميح (2) الفرق بين = و ==: العلامة (=) لتخزين قيمة في متغير، بينما (==) للمقارنة والتحقق من التساوي.",
            "• تلميح (3) المعامل المنطقي and: يشترط أن يتحقق الشرطان معاً (التفوق في المعدل + عدم الغياب) للحصول على المكافأة."
        ],
        bg_hex="FFFBEB",
        border_hex="D97706",
        title_color=(180, 83, 9)
    )

    skeleton_2 = [
        '# 1. قراءة العلامات الثلاث وعدد أيام الغياب',
        'g1 = int(input("أدخل علامة الاختبار الأول: "))',
        'g2 = int(input("أدخل علامة الاختبار الثاني: "))',
        'g3 = int(input("أدخل علامة الاختبار الثالث: "))',
        'absence = int(input("أدخل عدد أيام الغياب: "))',
        '',
        '# 2. حساب المتوسط الحسابي (انتبه للأقواس)',
        'avg = (g1 + g2 + g3) ........ 3',
        'print("معدلك الحسابي هو:", avg)',
        '',
        '# 3. تحديد التقدير باستخدام الشروط المتعددة if / elif / else',
        'if avg >= 90:',
        '    print("التقدير: ممتاز")',
        '............ avg >= 75:',
        '    print("التقدير: جيد جداً")',
        'elif avg ........ 50:',
        '    print("التقدير: ناجح")',
        'else:',
        '    print("التقدير: يحتاج إلى تحسين")',
        '',
        '# 4. شرط شهادة التفوق والانضباط باستخدام المعامل المنطقي (and)',
        'if avg >= 90 ............ absence ........ 0:',
        '    print("مبارك! تستحق شهادة التفوق ومكافأة 25 ديناراً أردنياً.")'
    ]
    add_code_box(doc, "💻 مساحة كتابة الكود (أكمل الفراغات البرمجية التالية):", skeleton_2, is_solution=False)

    add_rtl_paragraph(
        doc,
        "🚀 تحدي المبدعين (إضافي): أضف جملة شرطية تستخدم المعامل المنطقي (or) تطبع 'تنبيه من المرشد الأكاديمي!' إذا كان المعدل أقل من 50 أو كان عدد أيام الغياب أكبر من 5 أيام.",
        bold=True, size=10, color=(109, 40, 217), space_before=6, space_after=4
    )

    if include_solution:
        doc.add_page_break()
        add_header_banner(
            doc,
            ws_num_title="نموذج الإجابة المعتمد - ورقة عمل (2)",
            project_title="✅ الحل النموذجي: المرشد الأكاديمي وحاسبة المعدل الذكية",
            concepts_text="إشراف المعلم: أ. باسم النعيمي | مدرسة صويلح الثانوية للبنين | الصف الأول الثانوي الأكاديمي",
            theme_color_hex="15803D",
            light_bg_hex="F0FDF4"
        )
        sol_2 = [
            '# 1. قراءة العلامات الثلاث وعدد أيام الغياب',
            'g1 = int(input("أدخل علامة الاختبار الأول: "))',
            'g2 = int(input("أدخل علامة الاختبار الثاني: "))',
            'g3 = int(input("أدخل علامة الاختبار الثالث: "))',
            'absence = int(input("أدخل عدد أيام الغياب: "))',
            '',
            '# 2. حساب المتوسط الحسابي (انتبه للأقواس)',
            'avg = (g1 + g2 + g3) / 3',
            'print("معدلك الحسابي هو:", avg)',
            '',
            '# 3. تحديد التقدير باستخدام الشروط المتعددة if / elif / else',
            'if avg >= 90:',
            '    print("التقدير: ممتاز")',
            'elif avg >= 75:',
            '    print("التقدير: جيد جداً")',
            'elif avg >= 50:',
            '    print("التقدير: ناجح")',
            'else:',
            '    print("التقدير: يحتاج إلى تحسين")',
            '',
            '# 4. شرط شهادة التفوق والانضباط باستخدام المعامل المنطقي (and)',
            'if avg >= 90 and absence == 0:',
            '    print("مبارك! تستحق شهادة التفوق ومكافأة 25 ديناراً أردنياً.")',
            '',
            '# --- حل تحدي المبدعين الإضافي باستخدام (or) ---',
            'if avg < 50 or absence > 5:',
            '    print("تنبيه من المرشد الأكاديمي: يرجى مراجعة الإرشاد التربوي!")'
        ]
        add_code_box(doc, "🗝️ الكود البرمجي الكامل مع حل التحدي الإضافي:", sol_2, is_solution=True)


def build_ws3(doc, include_solution=False):
    add_header_banner(
        doc,
        ws_num_title="ورقة عمل (3) - مشروع مصغر شامل",
        project_title="🎢 مشروع: البوابة الذكية لمدينة الألعاب التفاعلية",
        concepts_text="باقي القسمة (%) | المعاملات المنطقية (and , or , not) | المقارنة (>= , < , ==) | الشروط المتداخلة if / else",
        theme_color_hex="6D28D9",
        light_bg_hex="F5F3FF"
    )

    add_box_section(
        doc,
        title="💡 فكرة المشروع (تطبيق شامل لكل ما تعلمته):",
        body_lines=[
            "كلّفتك إدارة مدينة الألعاب في عمّان ببرمجة البوابة الإلكترونية للعبة (قطار المغامرات السريع).",
            "يقوم البرنامج بالتأكد من شروط الأمان للزائر، ويحسب سعر تذكرته بالدينار الأردني، ثم يفحص رقم تذكرته لمعرفة إن كان من الفائزين بوجبة مجانية (أصحاب التذاكر ذات الأرقام الزوجية)!"
        ],
        bg_hex="F5F3FF",
        border_hex="7C3AED",
        title_color=(109, 40, 217)
    )

    add_rtl_paragraph(doc, "🛠️ قوانين البوابة الذكية (خطوات بناء المشروع):", bold=True, size=11.5, color=(15, 23, 42), space_before=8, space_after=4)
    steps = [
        "1) اطلب من الزائر إدخال: عمره (age)، وطوله بالسنتيمتر (height)، ورقم تذكرته (ticket_num).",
        "2) شرط الأمان لدخول القطار: يُسمح بالركوب إذا كان العمر 12 سنة أو أكثر (and) الطول 140 سم أو أكثر. غير ذلك يُطبع اعتذار بعدم السماح بالركوب.",
        "3) حساب سعر التذكرة بالدينار الأردني: السعر الأساسي هو (5 دنانير أردنية)، لكن إذا كان عمر الزائر أقل من 16 سنة أو (or) عمره 60 سنة فأكثر يحصل على تذكرة مخفضة بسعر (3 دنانير أردنية) فقط!",
        "4) هدية الحظ (باقي القسمة %): إذا كان باقي قسمة رقم التذكرة على 2 يساوي صفراً (ticket_num % 2 == 0)، فهذا يعني أن الرقم زوجي، فيطبع البرنامج: 'رقم تذكرتك زوجي! ربحت قسيمة شرائية بقيمة 2 دينار أردني'."
    ]
    for s in steps:
        add_rtl_paragraph(doc, s, size=10.5, color=(30, 41, 59), space_after=3)

    add_box_section(
        doc,
        title="🔑 تلميحات ذكية لمساعدتك (Hints):",
        body_lines=[
            "• تلميح (1) المعامل or: يكفي أن يتحقق شرط واحد فقط من الشرطين ليحصل الزائر على التذكرة المخفضة (3 دنانير).",
            "• تلميح (2) معامل باقي القسمة (%): يعطينا الباقي من القسمة؛ مثلاً 104 % 2 الناتج 0 (زوجي)، بينما 105 % 2 الناتج 1 (فردي).",
            "• تلميح (3) المعامل not: يُستخدم لعكس النتيجة المنطقية (تحويل True إلى False والعكس)."
        ],
        bg_hex="FFFBEB",
        border_hex="D97706",
        title_color=(180, 83, 9)
    )

    skeleton_3 = [
        '# 1. إدخال بيانات الزائر',
        'age = int(input("أدخل عمرك: "))',
        'height = int(input("أدخل طولك (سم): "))',
        'ticket_num = int(input("أدخل رقم تذكرتك: "))',
        '',
        '# 2. التحقق من شرط الأمان (العمر 12 فأكثر والطول 140 فأكثر معاً)',
        'if age >= 12 ............ height ........ 140:',
        '    print("مسموح لك بركوب قطار المغامرات!")',
        '',
        '    # 3. تحديد سعر التذكرة بالدينار الأردني (أقل من 16 أو 60 فأكثر)',
        '    if age < 16 ........ age >= 60:',
        '        price = 3',
        '        print("شملتك فئة الخصم! سعر تذكرتك:", price, "دنانير أردنية.")',
        '    else:',
        '        price = 5',
        '        print("سعر تذكرتك الأساسي:", price, "دنانير أردنية.")',
        '',
        '    # 4. فحص رقم التذكرة الزوجي باستخدام باقي القسمة (%)',
        '    if ticket_num ........ 2 == 0:',
        '        print("رقم تذكرتك زوجي! ربحت قسيمة مجانية بقيمة 2 دينار أردني.")',
        'else:',
        '    print("عذراً، شروط الأمان (العمر أو الطول) غير مستوفاة.")'
    ]
    add_code_box(doc, "💻 مساحة كتابة الكود (أكمل الفراغات البرمجية التالية):", skeleton_3, is_solution=False)

    add_rtl_paragraph(
        doc,
        "🚀 تحدي المبدعين (إضافي): أضف سطراً يحسب (نقاط المكافآت) برفع عدد الألعاب التي لعبها الزائر للأس 2 باستخدام معامل الأس (**)، وجرّب استخدام (not) في شرط إضافي!",
        bold=True, size=10, color=(109, 40, 217), space_before=6, space_after=4
    )

    if include_solution:
        doc.add_page_break()
        add_header_banner(
            doc,
            ws_num_title="نموذج الإجابة المعتمد - ورقة عمل (3)",
            project_title="✅ الحل النموذجي: البوابة الذكية لمدينة الألعاب التفاعلية",
            concepts_text="إشراف المعلم: أ. باسم النعيمي | مدرسة صويلح الثانوية للبنين | الصف الأول الثانوي الأكاديمي",
            theme_color_hex="15803D",
            light_bg_hex="F0FDF4"
        )
        sol_3 = [
            '# 1. إدخال بيانات الزائر',
            'age = int(input("أدخل عمرك: "))',
            'height = int(input("أدخل طولك (سم): "))',
            'ticket_num = int(input("أدخل رقم تذكرتك: "))',
            '',
            '# 2. التحقق من شرط الأمان (العمر 12 فأكثر والطول 140 فأكثر معاً)',
            'if age >= 12 and height >= 140:',
            '    print("مسموح لك بركوب قطار المغامرات!")',
            '',
            '    # 3. تحديد سعر التذكرة بالدينار الأردني (أقل من 16 أو 60 فأكثر)',
            '    if age < 16 or age >= 60:',
            '        price = 3',
            '        print("شملتك فئة الخصم! سعر تذكرتك:", price, "دنانير أردنية.")',
            '    else:',
            '        price = 5',
            '        print("سعر تذكرتك الأساسي:", price, "دنانير أردنية.")',
            '',
            '    # 4. فحص رقم التذكرة الزوجي باستخدام باقي القسمة (%)',
            '    if ticket_num % 2 == 0:',
            '        print("رقم تذكرتك زوجي! ربحت قسيمة مجانية بقيمة 2 دينار أردني.")',
            'else:',
            '    print("عذراً، شروط الأمان (العمر أو الطول) غير مستوفاة.")',
            '',
            '# --- حل تحدي المبدعين الإضافي باستخدام الأس (**) والمعامل (not) ---',
            'games_played = 4',
            'points = games_played ** 2',
            'if not (age < 12):',
            '    print("نقاط المكافآت الخاصة بك:", points)'
        ]
        add_code_box(doc, "🗝️ الكود البرمجي الكامل مع حل التحدي الإضافي:", sol_3, is_solution=True)


def configure_doc_margins(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)


def main():
    os.makedirs("word-files", exist_ok=True)

    # 1. All-in-one master file (with solutions)
    doc_all = Document()
    configure_doc_margins(doc_all)
    build_ws1(doc_all, include_solution=True)
    doc_all.add_page_break()
    build_ws2(doc_all, include_solution=True)
    doc_all.add_page_break()
    build_ws3(doc_all, include_solution=True)
    master_filename = "أوراق_عمل_بايثون_الأول_الثانوي_الأكاديمي_مع_الحلول.docx"
    doc_all.save(master_filename)
    doc_all.save(os.path.join("word-files", master_filename))

    # 2. Student version (all 3 worksheets without solutions)
    doc_student = Document()
    configure_doc_margins(doc_student)
    build_ws1(doc_student, include_solution=False)
    doc_student.add_page_break()
    build_ws2(doc_student, include_solution=False)
    doc_student.add_page_break()
    build_ws3(doc_student, include_solution=False)
    student_filename = "نسخة_الطالب_أوراق_عمل_بايثون_بدون_حلول.docx"
    doc_student.save(student_filename)
    doc_student.save(os.path.join("word-files", student_filename))

    # 3. Individual worksheet files (each with its solution)
    doc_1 = Document()
    configure_doc_margins(doc_1)
    build_ws1(doc_1, include_solution=True)
    f1 = "ورقة_عمل_1_الكاشير_الذكي_مع_الحل.docx"
    doc_1.save(f1)
    doc_1.save(os.path.join("word-files", f1))

    doc_2 = Document()
    configure_doc_margins(doc_2)
    build_ws2(doc_2, include_solution=True)
    f2 = "ورقة_عمل_2_المرشد_الأكاديمي_مع_الحل.docx"
    doc_2.save(f2)
    doc_2.save(os.path.join("word-files", f2))

    doc_3 = Document()
    configure_doc_margins(doc_3)
    build_ws3(doc_3, include_solution=True)
    f3 = "ورقة_عمل_3_بوابة_مدينة_الألعاب_مع_الحل.docx"
    doc_3.save(f3)
    doc_3.save(os.path.join("word-files", f3))

    print("Successfully generated all 5 Word (.docx) files!")


if __name__ == "__main__":
    main()
